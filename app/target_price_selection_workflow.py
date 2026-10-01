from __future__ import annotations

"""Target Price workflow: Excel-based source selection.

Flow:
1) ordinary import -> Calculate Cost -> calculated Excel with all options;
2) user fills "Источник для расчета таргета" in that same workbook;
3) the same Import button restores a self-contained batch from the workbook, including after restart;
4) "Расчет Target" validates the restored DB batch and saves the reverse target-price calculation.
"""

from collections import defaultdict
from datetime import datetime
from decimal import Decimal
import logging
from pathlib import Path
import re
from typing import Any

from app.utils.text import clean_multi_spaces

logger = logging.getLogger(__name__)

SELECTION_HEADER = "Источник для расчета таргета"
_SELECTION_HEADER_NORM = "".join(ch for ch in SELECTION_HEADER.lower() if ch.isalnum())
_TECH_SHEET = "_TargetPriceOptions"
_installed = False


def _norm_header(value: object) -> str:
    text = clean_multi_spaces(str(value or "")).lower()
    return "".join(ch for ch in text if ch.isalnum())


def _norm_text(value: object) -> str:
    return clean_multi_spaces(str(value or "")).casefold()


def _norm_article(value: object) -> str:
    text = clean_multi_spaces(str(value or ""))
    if text.endswith(".0"):
        try:
            return str(int(float(text)))
        except Exception:
            pass
    return text.casefold()


def _is_selection_workbook(file_path: str | Path) -> bool:
    if Path(file_path).suffix.lower() == ".xls":
        return False
    from openpyxl import load_workbook

    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        # A prepared Target workbook is defined by its content, not by which
        # sheet happened to be active when the user saved Excel.
        worksheets = []
        if "Calculated" in wb.sheetnames:
            worksheets.append(wb["Calculated"])
        worksheets.extend(
            ws for ws in wb.worksheets
            if ws.title != "Calculated" and ws.title != _TECH_SHEET
        )
        for ws in worksheets:
            header = next(
                ws.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            ) or ()
            if any(
                _norm_header(value) == _SELECTION_HEADER_NORM
                for value in header
            ):
                return True
        return False
    finally:
        wb.close()


def _decimal_or_none(value: object) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, Decimal):
        return value
    try:
        text = str(value).strip().replace(" ", "").replace("\xa0", "").replace(",", ".")
        if not text or text == "-":
            return None
        return Decimal(text)
    except Exception:
        return None


def _decimal_or_zero(value: object) -> Decimal:
    return _decimal_or_none(value) or Decimal("0")


def _datetime_or_none(value: object):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    text = clean_multi_spaces(str(value))
    if not text or text == "-":
        return None
    for fmt in (
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text)
    except Exception:
        return None


def _read_selection_workbook(file_path: str | Path) -> dict[str, Any]:
    """Read a calculated workbook without relying on any live temp batch.

    New workbooks also contain a very-hidden technical sheet with the exact
    option metadata.  Older workbooks (already created before this fix) are
    still supported from the visible Supplier_N / Full Cost Msk_N columns.
    """
    from openpyxl import load_workbook

    aliases = {
        "supplierarticle": "supplier_article",
        "supplierproductname": "supplier_product_name",
        "ourproductname": "our_product_name",
        _SELECTION_HEADER_NORM: "target_source",
    }

    wb = load_workbook(file_path, data_only=True, read_only=True)
    try:
        ws = wb["Calculated"] if "Calculated" in wb.sheetnames else wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            return {"rows": [], "tech_options": []}

        columns: dict[int, str] = {}
        dynamic: dict[int, dict[str, int]] = defaultdict(dict)
        for index, value in enumerate(header):
            raw = clean_multi_spaces(str(value or ""))
            key = aliases.get(_norm_header(raw))
            if key:
                columns[index] = key
                continue
            match = re.fullmatch(
                r"(Cost Novo with VAT|Full Cost Msk|Supplier|last update|Currency)_(\d+)",
                raw,
                flags=re.IGNORECASE,
            )
            if not match:
                continue
            field_label = match.group(1).casefold()
            option_index = int(match.group(2))
            field_name = {
                "cost novo with vat": "cost_novo_wvat",
                "full cost msk": "full_cost_msk",
                "supplier": "supplier_name",
                "last update": "price_date_used",
                "currency": "currency_code",
            }[field_label]
            dynamic[option_index][field_name] = index

        if "target_source" not in columns.values():
            return {"rows": [], "tech_options": []}
        if "our_product_name" not in columns.values():
            raise ValueError("В файле выбора должна быть колонка 'Our Product Name'.")

        rows: list[dict[str, Any]] = []
        for excel_row, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            row: dict[str, Any] = {
                "excel_row": excel_row,
                "supplier_article": "",
                "supplier_product_name": "",
                "our_product_name": "",
                "target_source": "",
                "options": [],
            }
            for index, key in columns.items():
                value = values[index] if index < len(values) else None
                row[key] = clean_multi_spaces(str(value or ""))

            for option_index in sorted(dynamic):
                field_indexes = dynamic[option_index]
                supplier_idx = field_indexes.get("supplier_name")
                full_cost_idx = field_indexes.get("full_cost_msk")
                supplier_name = (
                    clean_multi_spaces(str(values[supplier_idx] or ""))
                    if supplier_idx is not None and supplier_idx < len(values)
                    else ""
                )
                full_cost = (
                    values[full_cost_idx]
                    if full_cost_idx is not None and full_cost_idx < len(values)
                    else None
                )
                if not supplier_name and _decimal_or_none(full_cost) is None:
                    continue
                option: dict[str, Any] = {"opt_rank": option_index}
                for field_name, source_index in field_indexes.items():
                    value = values[source_index] if source_index < len(values) else None
                    option[field_name] = value
                row["options"].append(option)

            if row["our_product_name"] or row["target_source"] or row["options"]:
                rows.append(row)

        tech_options: list[dict[str, Any]] = []
        if _TECH_SHEET in wb.sheetnames:
            tech = wb[_TECH_SHEET]
            tech_header = next(
                tech.iter_rows(min_row=1, max_row=1, values_only=True),
                None,
            ) or ()
            tech_names = [clean_multi_spaces(str(value or "")) for value in tech_header]
            for values in tech.iter_rows(min_row=2, values_only=True):
                record = {
                    tech_names[index]: (values[index] if index < len(values) else None)
                    for index in range(len(tech_names))
                    if tech_names[index]
                }
                if record:
                    tech_options.append(record)

        return {"rows": rows, "tech_options": tech_options}
    finally:
        wb.close()


def _read_selection_rows(file_path: str | Path) -> list[dict[str, Any]]:
    # Backwards-compatible helper name used by older local code.
    return list(_read_selection_workbook(file_path).get("rows") or [])

def _install_calculated_export_source_column() -> None:
    from app.db.models import Product, TempTargetPriceImport, TempTargetPriceOption
    from app.exports.target_price_exporter import TargetPriceExporter

    current = TargetPriceExporter.export_calculated
    if getattr(current, "_target_source_column", False):
        return

    def export_calculated_with_source(
        self,
        batch_id: str,
        imported_by: str,
        file_path: str | Path,
    ) -> Path:
        rows, max_opt = self._collect_calculated_rows(batch_id, imported_by)

        temp_rows = (
            self.session.query(TempTargetPriceImport)
            .filter(
                TempTargetPriceImport.batch_id == batch_id,
                TempTargetPriceImport.imported_by == imported_by,
            )
            .order_by(
                TempTargetPriceImport.import_row_no.asc(),
                TempTargetPriceImport.id.asc(),
            )
            .all()
        )
        selected_ids = {
            int(row.selected_option_id)
            for row in temp_rows
            if row.selected_option_id is not None
        }
        selected_names = {
            int(option.id): option.supplier_name or ""
            for option in (
                self.session.query(TempTargetPriceOption)
                .filter(TempTargetPriceOption.id.in_(selected_ids))
                .all()
                if selected_ids
                else []
            )
        }

        for data, temp_row in zip(rows, temp_rows):
            data[SELECTION_HEADER] = (
                selected_names.get(int(temp_row.selected_option_id), "")
                if temp_row.selected_option_id is not None
                else ""
            )

        base_headers = [
            "Supplier Article",
            "Supplier Product Name",
            "Our Product Name",
            SELECTION_HEADER,
            "Дистр цена",
            "Промо цена",
            "curr LPC",
            "curr Landed cost",
        ]
        dynamic_headers: list[str] = []
        for index in range(1, max_opt + 1):
            dynamic_headers.extend(
                [
                    f"Cost Novo with VAT_{index}",
                    f"Full Cost Msk_{index}",
                    f"Supplier_{index}",
                    f"last update_{index}",
                    f"Currency_{index}",
                ]
            )

        output = self._save_workbook(
            file_path,
            "Calculated",
            base_headers + dynamic_headers,
            rows,
            freeze_cell="E2",
            format_mode="target_price_calculated",
        )

        # The visible sheet stays exactly the same apart from the one source
        # column requested by the user.  Exact option metadata is stored on a
        # very-hidden technical sheet so the workbook is self-contained and can
        # be imported after an application restart or for another target supplier.
        product_ids = {
            int(row.selected_product_id)
            for row in temp_rows
            if row.selected_product_id is not None
        }
        product_names = {
            int(product.id): product.name or ""
            for product in (
                self.session.query(Product).filter(Product.id.in_(product_ids)).all()
                if product_ids
                else []
            )
        }
        options = (
            self.session.query(TempTargetPriceOption)
            .filter(
                TempTargetPriceOption.batch_id == batch_id,
                TempTargetPriceOption.imported_by == imported_by,
            )
            .order_by(
                TempTargetPriceOption.temp_import_id.asc(),
                TempTargetPriceOption.opt_rank.asc(),
                TempTargetPriceOption.id.asc(),
            )
            .all()
        )
        temp_by_id = {int(row.id): row for row in temp_rows}

        tech_headers = [
            "Supplier Article",
            "Supplier Product Name",
            "Our Product Name",
            "opt_rank",
            "supplier_id",
            "source_type",
            "supplier_name",
            "supplier_price",
            "price_date_used",
            "cost_novo_wvat",
            "full_cost_msk",
            "currency_code",
            "fx_rate_used",
            "fx_markup_used",
            "fx_markup_abs_used",
            "transport_used",
            "reexport_used",
            "insurance_used",
            "agent_fee_used",
            "has_customs_used",
            "via_novo_used",
            "bank_fee_used",
            "customs_fee_used",
            "move_used",
            "is_excise_used",
            "additional_customs_used",
            "storage_used",
            "marking_used",
        ]
        tech_rows: list[list[Any]] = []
        for option in options:
            temp_row = temp_by_id.get(int(option.temp_import_id))
            if temp_row is None:
                continue
            product_name = (
                product_names.get(int(temp_row.selected_product_id), "")
                if temp_row.selected_product_id is not None
                else ""
            )
            tech_rows.append(
                [
                    temp_row.supplier_article,
                    temp_row.product_name,
                    product_name,
                    option.opt_rank,
                    option.supplier_id,
                    option.source_type,
                    option.supplier_name,
                    option.supplier_price,
                    option.price_date_used,
                    option.cost_novo_wvat,
                    option.full_cost_msk,
                    option.currency_code,
                    option.fx_rate_used,
                    option.fx_markup_used,
                    option.fx_markup_abs_used,
                    option.transport_used,
                    option.reexport_used,
                    option.insurance_used,
                    option.agent_fee_used,
                    option.has_customs_used,
                    option.via_novo_used,
                    option.bank_fee_used,
                    option.customs_fee_used,
                    option.move_used,
                    option.is_excise_used,
                    option.additional_customs_used,
                    option.storage_used,
                    option.marking_used,
                ]
            )

        if tech_rows:
            # Add the technical sheet through Excel COM as well.  Re-saving a
            # COM-created workbook through openpyxl is intentionally avoided.
            import pythoncom

            excel = None
            workbook = None
            try:
                excel = self._create_excel_app()
                workbook = excel.Workbooks.Open(str(Path(output).resolve()))
                existing = None
                for sheet_index in range(1, workbook.Worksheets.Count + 1):
                    candidate = workbook.Worksheets(sheet_index)
                    if candidate.Name == _TECH_SHEET:
                        existing = candidate
                        break
                if existing is not None:
                    existing.Delete()
                tech = workbook.Worksheets.Add(After=workbook.Worksheets(workbook.Worksheets.Count))
                tech.Name = _TECH_SHEET

                def excel_safe(value):
                    if isinstance(value, Decimal):
                        return float(value)
                    return value

                matrix = [tech_headers] + tech_rows
                matrix = tuple(tuple(excel_safe(value) for value in row) for row in matrix)
                target = tech.Range(
                    tech.Cells(1, 1),
                    tech.Cells(len(matrix), len(tech_headers)),
                )
                target.Value = matrix
                tech.Visible = 2  # xlSheetVeryHidden
                workbook.Save()
            finally:
                try:
                    if workbook is not None:
                        workbook.Close(SaveChanges=False)
                except Exception:
                    logger.exception("Не удалось закрыть Target Price workbook")
                try:
                    if excel is not None:
                        excel.Quit()
                except Exception:
                    logger.exception("Не удалось закрыть Excel после Target Price export")
                try:
                    pythoncom.CoUninitialize()
                except Exception:
                    pass

        return output

    export_calculated_with_source._target_source_column = True
    export_calculated_with_source._target_source_original = current
    TargetPriceExporter.export_calculated = export_calculated_with_source

def _apply_selection_rows(
    session,
    *,
    batch_id: str,
    imported_by: str,
    selection_rows: list[dict[str, Any]],
    tech_options: list[dict[str, Any]] | None = None,
) -> dict[str, int]:
    """Restore temp Target Price rows/options entirely from the workbook.

    The workbook is intentionally independent of the previous application
    session and target supplier.  This lets the same selected-source file be
    reused after restart and for different target suppliers with identical
    target calculation conditions.
    """
    from app.db.models import Product, Supplier, TempTargetPriceImport, TempTargetPriceOption
    from app.services.cost_calculation_service import CostCalculationService
    from app.services.product_uc3_service import (
        SOURCE_MANUAL,
        SOURCE_SUPPLIER,
        SOURCE_TARGET_UC3,
        SOURCE_WALK_AWAY_UC3,
    )
    from app.utils.text import normalize_product_name

    if not selection_rows:
        raise ValueError("В файле нет строк для импорта выбора.")

    # This import creates a fresh, self-contained temp batch.  No dependency on
    # the old batch is allowed: it may have been deleted or belong to a previous
    # program session.
    session.query(TempTargetPriceOption).filter(
        TempTargetPriceOption.batch_id == batch_id,
        TempTargetPriceOption.imported_by == imported_by,
    ).delete(synchronize_session=False)
    session.query(TempTargetPriceImport).filter(
        TempTargetPriceImport.batch_id == batch_id,
        TempTargetPriceImport.imported_by == imported_by,
    ).delete(synchronize_session=False)
    session.flush()

    requested_names = {
        clean_multi_spaces(row.get("our_product_name"))
        for row in selection_rows
        if clean_multi_spaces(row.get("our_product_name"))
    }
    exact_products = {
        str(product.name).strip(): product
        for product in (
            session.query(Product).filter(Product.name.in_(requested_names)).all()
            if requested_names
            else []
        )
        if product.name
    }
    missing_names = requested_names - set(exact_products)
    normalized_products: dict[str, Product] = {}
    if missing_names:
        # Fallback is only for renamed spacing/case variants.  Keep ambiguous
        # normalized names unresolved instead of silently choosing one.
        duplicates: set[str] = set()
        for product in session.query(Product).filter(Product.name.isnot(None)).all():
            key = normalize_product_name(product.name)
            if not key:
                continue
            if key in normalized_products and normalized_products[key].id != product.id:
                duplicates.add(key)
            else:
                normalized_products[key] = product
        for key in duplicates:
            normalized_products.pop(key, None)

    def resolve_product(name: str, excel_row: int) -> Product:
        product = exact_products.get(name)
        if product is None:
            key = normalize_product_name(name)
            product = normalized_products.get(key) if key else None
        if product is None:
            raise ValueError(
                f"Строка Excel {excel_row}: продукт '{name}' не найден в Products."
            )
        return product

    suppliers = session.query(Supplier).all()
    suppliers_by_id = {int(supplier.id): supplier for supplier in suppliers}
    suppliers_by_name = {
        _norm_text(supplier.name): supplier
        for supplier in suppliers
        if supplier.name
    }

    def source_type_for(name: object, explicit: object = None) -> str:
        explicit_text = clean_multi_spaces(str(explicit or ""))
        if explicit_text:
            return explicit_text
        normalized = _norm_text(name)
        if normalized == _norm_text("Target uC3"):
            return SOURCE_TARGET_UC3
        if normalized == _norm_text("Walk-Away uC3"):
            return SOURCE_WALK_AWAY_UC3
        if normalized == _norm_text("Manual"):
            return SOURCE_MANUAL
        return SOURCE_SUPPLIER

    def bool_value(value: object, default: bool = False) -> bool:
        if value is None or value == "":
            return default
        if isinstance(value, bool):
            return value
        return clean_multi_spaces(str(value)).casefold() in {
            "1", "true", "yes", "да", "y",
        }

    def identity_from_values(our_name: object, article: object, supplier_product: object):
        return (
            _norm_text(our_name),
            _norm_article(article),
            _norm_text(supplier_product),
        )

    tech_by_identity: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for record in tech_options or []:
        key = identity_from_values(
            record.get("Our Product Name"),
            record.get("Supplier Article"),
            record.get("Supplier Product Name"),
        )
        if key[0]:
            tech_by_identity[key].append(record)
    for records in tech_by_identity.values():
        records.sort(key=lambda item: int(_decimal_or_zero(item.get("opt_rank"))))

    calc_service = CostCalculationService(session)
    fixed = calc_service.get_fixed_costs()
    zero = Decimal("0")
    fixed_bank = _decimal_or_zero(getattr(fixed, "bank_fee", 0))
    fixed_customs = _decimal_or_zero(getattr(fixed, "customs_fee", 0))
    fixed_add_customs = _decimal_or_zero(getattr(fixed, "additional_customs", 0))
    fixed_storage = _decimal_or_zero(getattr(fixed, "storage", 0))
    fixed_move = _decimal_or_zero(getattr(fixed, "move", 0))
    marking_by_product: dict[int, Decimal] = {}

    def marking_for(product_id: int) -> Decimal:
        product_id = int(product_id)
        if product_id not in marking_by_product:
            marking_by_product[product_id] = _decimal_or_zero(
                calc_service.get_marking_cost(product_id)
            )
        return marking_by_product[product_id]

    restored_rows: list[tuple[TempTargetPriceImport, dict[str, Any], Product]] = []
    for sequence, source in enumerate(selection_rows, start=1):
        excel_row = int(source.get("excel_row") or sequence + 1)
        product_name = clean_multi_spaces(source.get("our_product_name"))
        if not product_name:
            if clean_multi_spaces(source.get("target_source")):
                raise ValueError(
                    f"Строка Excel {excel_row}: заполнен '{SELECTION_HEADER}', "
                    "но пустой 'Our Product Name'."
                )
            continue
        product = resolve_product(product_name, excel_row)
        temp = TempTargetPriceImport(
            batch_id=batch_id,
            imported_by=imported_by,
            import_row_no=sequence,
            import_date=datetime.utcnow(),
            target_supplier_id=None,
            supplier_article=clean_multi_spaces(source.get("supplier_article")) or None,
            product_name=clean_multi_spaces(source.get("supplier_product_name")) or None,
            selected_product_id=int(product.id),
            selected_option_id=None,
            new_product_name=None,
            new_brand=None,
            new_pack=None,
            new_qty_in_box=None,
            new_is_excise=None,
        )
        session.add(temp)
        restored_rows.append((temp, source, product))
    session.flush()

    selected_count = 0
    blank_count = 0
    unavailable_count = 0
    option_count = 0

    for temp, source, product in restored_rows:
        identity = identity_from_values(
            product.name,
            temp.supplier_article,
            temp.product_name,
        )
        source_options = tech_by_identity.get(identity)
        if source_options:
            option_records = list(source_options)
        else:
            option_records = list(source.get("options") or [])

        if not option_records:
            # A product can legitimately have no usable cost source.  Do not
            # abort the whole prepared-file import because of one such row.
            # Keep the product row in staging without selected_option_id; the
            # user can later recalculate/add a source for it independently.
            unavailable_count += 1
            continue

        created_options: list[TempTargetPriceOption] = []
        for fallback_rank, data in enumerate(option_records, start=1):
            supplier_name = clean_multi_spaces(
                data.get("supplier_name")
                if "supplier_name" in data
                else data.get("supplier_name", "")
            )
            # Technical sheet stores supplier_name under the model field name;
            # visible-sheet fallback uses the same normalized key.
            if not supplier_name:
                supplier_name = clean_multi_spaces(data.get("Supplier"))
            if not supplier_name:
                continue

            source_type = source_type_for(supplier_name, data.get("source_type"))
            supplier_id = None
            raw_supplier_id = _decimal_or_none(data.get("supplier_id"))
            if raw_supplier_id is not None:
                candidate = suppliers_by_id.get(int(raw_supplier_id))
                if candidate is not None:
                    supplier_id = int(candidate.id)
            if supplier_id is None and source_type == SOURCE_SUPPLIER:
                supplier = suppliers_by_name.get(_norm_text(supplier_name))
                if supplier is not None:
                    supplier_id = int(supplier.id)

            full_cost = _decimal_or_none(data.get("full_cost_msk"))
            cost_novo = _decimal_or_none(data.get("cost_novo_wvat"))
            if full_cost is None:
                continue

            has_tech = "fx_rate_used" in data or "bank_fee_used" in data
            option = TempTargetPriceOption(
                temp_import_id=int(temp.id),
                batch_id=batch_id,
                imported_by=imported_by,
                calc_date=datetime.utcnow(),
                supplier_id=supplier_id,
                source_type=source_type,
                product_id=int(product.id),
                supplier_name=supplier_name,
                supplier_article=temp.supplier_article,
                supplier_product_name=temp.product_name,
                supplier_price=_decimal_or_zero(data.get("supplier_price")),
                price_date_used=_datetime_or_none(data.get("price_date_used")),
                cost_novo_wvat=cost_novo or zero,
                full_cost_msk=full_cost,
                currency_code=clean_multi_spaces(data.get("currency_code")) or "-",
                fx_rate_used=_decimal_or_zero(data.get("fx_rate_used")),
                fx_markup_used=_decimal_or_zero(data.get("fx_markup_used")),
                fx_markup_abs_used=_decimal_or_zero(data.get("fx_markup_abs_used")),
                transport_used=_decimal_or_zero(data.get("transport_used")),
                reexport_used=_decimal_or_zero(data.get("reexport_used")),
                insurance_used=_decimal_or_zero(data.get("insurance_used")),
                agent_fee_used=_decimal_or_zero(data.get("agent_fee_used")),
                has_customs_used=bool_value(data.get("has_customs_used"), False),
                via_novo_used=bool_value(data.get("via_novo_used"), False),
                bank_fee_used=(
                    _decimal_or_zero(data.get("bank_fee_used")) if has_tech else fixed_bank
                ),
                customs_fee_used=(
                    _decimal_or_zero(data.get("customs_fee_used")) if has_tech else fixed_customs
                ),
                move_used=(
                    _decimal_or_zero(data.get("move_used")) if has_tech else fixed_move
                ),
                is_excise_used=bool_value(
                    data.get("is_excise_used"),
                    bool(getattr(product, "is_excise", False)),
                ),
                additional_customs_used=(
                    _decimal_or_zero(data.get("additional_customs_used"))
                    if has_tech else fixed_add_customs
                ),
                storage_used=(
                    _decimal_or_zero(data.get("storage_used")) if has_tech else fixed_storage
                ),
                marking_used=(
                    _decimal_or_zero(data.get("marking_used"))
                    if has_tech else marking_for(int(product.id))
                ),
                opt_rank=int(_decimal_or_zero(data.get("opt_rank"))) or fallback_rank,
            )
            session.add(option)
            created_options.append(option)
            option_count += 1

        session.flush()
        if not created_options:
            # The workbook contained option columns, but none had a usable
            # Full Cost Msk. Treat this the same as "no calculated options":
            # keep importing the rest of the file.
            unavailable_count += 1
            continue

        selected_name = clean_multi_spaces(source.get("target_source"))
        if not selected_name:
            blank_count += 1
            continue
        selected_norm = _norm_text(selected_name)
        selected = [
            option for option in created_options
            if _norm_text(option.supplier_name) == selected_norm
        ]
        if not selected:
            available = ", ".join(option.supplier_name for option in created_options)
            raise ValueError(
                f"Строка Excel {source.get('excel_row')}: источник '{selected_name}' "
                f"не найден для '{product.name}'. Доступно: {available}."
            )
        if len(selected) > 1:
            raise ValueError(
                f"Строка Excel {source.get('excel_row')}: источник '{selected_name}' "
                "неоднозначен. Выберите уникальный вариант."
            )
        temp.selected_option_id = int(selected[0].id)
        selected_count += 1

    session.flush()
    return {
        "rows_in_file": len(restored_rows),
        "matched_rows": len(restored_rows),
        "options": option_count,
        "selected": selected_count,
        "blank": blank_count,
        "unavailable": unavailable_count,
    }

def _install_target_supplier_fields_behavior() -> None:
    """Mirror Supplier Prices new-supplier name/country behavior."""
    from app.page_functions.target_prices_page import TargetPricesPage

    current_toggle = TargetPricesPage.toggle_new_supplier_field
    if getattr(current_toggle, "_target_supplier_country_behavior", False):
        return

    def toggle_new_supplier_field(self, enabled: bool):
        widgets = (
            self.ui.label_blank_new_supplier,
            self.ui.line_NewSupplier,
            self.ui.label_new_supplier_country,
            self.ui.line_NewSupplierCountry,
        )
        for widget in widgets:
            widget.setVisible(bool(enabled))
        self.ui.line_NewSupplier.setEnabled(bool(enabled))
        self.ui.line_NewSupplierCountry.setEnabled(bool(enabled))
        if enabled:
            self.ui.line_NewSupplier.setFocus()

    toggle_new_supplier_field._target_supplier_country_behavior = True
    toggle_new_supplier_field._target_supplier_country_original = current_toggle
    TargetPricesPage.toggle_new_supplier_field = toggle_new_supplier_field

    current_defaults = TargetPricesPage.apply_default_values

    def apply_default_values_with_country(self):
        current_defaults(self)
        if hasattr(self.ui, "line_NewSupplierCountry"):
            self.ui.line_NewSupplierCountry.clear()
        self.toggle_new_supplier_field(False)

    TargetPricesPage.apply_default_values = apply_default_values_with_country

    current_toggled = TargetPricesPage.on_new_supplier_toggled

    def on_new_supplier_toggled_with_country(self, checked: bool):
        if checked and hasattr(self.ui, "line_NewSupplierCountry"):
            self.ui.line_NewSupplierCountry.clear()
        current_toggled(self, checked)

    TargetPricesPage.on_new_supplier_toggled = on_new_supplier_toggled_with_country

    current_changed = TargetPricesPage.on_supplier_changed

    def on_supplier_changed_with_country(self):
        current_changed(self)
        supplier_id = self.ui.cbo_SupplName.currentData()
        if supplier_id is None or not hasattr(self.ui, "line_NewSupplierCountry"):
            return
        try:
            from app.services.supplier_service import SupplierService

            with self.get_session() as session:
                data = SupplierService(session).load_supplier_snapshot(int(supplier_id))
            self.ui.line_NewSupplierCountry.setText(data.country or "")
        except Exception:
            logger.exception("Не удалось загрузить страну поставщика в Target Price")

    TargetPricesPage.on_supplier_changed = on_supplier_changed_with_country

    current_form_data = TargetPricesPage.get_supplier_form_data

    def get_supplier_form_data_with_country(self):
        data = current_form_data(self)
        country = clean_multi_spaces(self.ui.line_NewSupplierCountry.text())
        if self.ui.cbx_NewSupplier.isChecked() and not country:
            raise ValueError("Введите страну нового поставщика.")
        # SupplierService.update_supplier intentionally preserves the existing
        # country when None is passed, matching Supplier Prices behavior.
        data.country = country or None
        return data

    TargetPricesPage.get_supplier_form_data = get_supplier_form_data_with_country


def _install_selection_aware_import() -> None:
    from app.utils import page_background_integration as pbi

    current_attach = pbi._attach_target_prices
    if getattr(current_attach, "_target_source_import", False):
        return

    def attach_target_prices_with_selection(page) -> None:
        if getattr(page, "_target_price_selection_workflow_attached", False):
            return
        page._target_price_selection_workflow_attached = True

        import copy
        from PySide6.QtCore import QUrl
        from PySide6.QtGui import QDesktopServices
        from PySide6.QtWidgets import QFileDialog

        if not hasattr(page, "_pending_deletes"):
            page._pending_deletes = set()
        if not hasattr(page, "_deleted_row_snapshots"):
            page._deleted_row_snapshots = []

        _install_calculated_export_source_column()
        current_attach(page)

        # In the new Excel-selection workflow "final Supplier" is no longer a
        # GUI choice.  The first calculated-mode column is the source imported
        # from the workbook, so name it exactly as in Excel.
        if getattr(page, "calc_headers", None):
            page.calc_headers[0] = SELECTION_HEADER

        old_import_slot = getattr(page, "_background_target_price_import_slot", None)
        if old_import_slot is not None:
            pbi._disconnect(page.ui.btn_Import.clicked, old_import_slot)

        old_calc_slot = getattr(page, "_background_target_price_calc_slot", None)
        if old_calc_slot is not None:
            pbi._disconnect(page.ui.btn_CalcCost.clicked, old_calc_slot)

        # The final Target calculation must not depend on a GUI-only flag such
        # as _showing_options.  A self-contained selection workbook can be
        # imported after an application restart and reused for another target
        # supplier, so readiness is determined from the restored temp batch.
        old_save_slot = getattr(page, "_background_target_price_save_slot", None)
        if old_save_slot is not None:
            pbi._disconnect(page.ui.btn_Save.clicked, old_save_slot)

        def _delete_pending(session, pending_deletes, batch_id, imported_by):
            if not pending_deletes:
                return
            from app.db.models import TempTargetPriceImport, TempTargetPriceOption

            session.query(TempTargetPriceOption).filter(
                TempTargetPriceOption.temp_import_id.in_(pending_deletes),
                TempTargetPriceOption.batch_id == batch_id,
                TempTargetPriceOption.imported_by == imported_by,
            ).delete(synchronize_session=False)
            session.query(TempTargetPriceImport).filter(
                TempTargetPriceImport.id.in_(pending_deletes),
                TempTargetPriceImport.batch_id == batch_id,
                TempTargetPriceImport.imported_by == imported_by,
            ).delete(synchronize_session=False)
            session.flush()

        def import_selection_aware() -> None:
            file_path, _ = QFileDialog.getOpenFileName(
                page,
                "Выберите файл",
                str(Path(__file__).resolve().parents[1]),
                "Excel files (*.xls *.xlsx)",
            )
            if not file_path:
                return

            try:
                is_selection = _is_selection_workbook(file_path)
            except Exception as exc:
                page.show_error_message(f"Не удалось прочитать заголовок Excel:\n{exc}")
                return

            # A returned selection workbook is self-contained.  Always create a
            # fresh batch from it, even after an application restart.  It is not
            # tied to the supplier currently selected on the page.
            page.start_new_batch()
            batch_id = page.batch_id
            imported_by = page.imported_by

            if is_selection:
                def work_selection(progress):
                    from app.db.db import SessionLocal

                    try:
                        progress("Target Price: импортирую выбранные источники...")
                        workbook_data = _read_selection_workbook(file_path)
                        selection_rows = list(workbook_data.get("rows") or [])
                        if not selection_rows:
                            raise ValueError("В файле нет строк для импорта выбора.")
                        with SessionLocal() as session:
                            result = _apply_selection_rows(
                                session,
                                batch_id=batch_id,
                                imported_by=imported_by,
                                selection_rows=selection_rows,
                                tech_options=list(workbook_data.get("tech_options") or []),
                            )
                            session.commit()
                        return result
                    finally:
                        SessionLocal.remove()

                def done_selection(payload):
                    from PySide6.QtCore import QTimer

                    page._showing_options = True
                    if getattr(page, "calc_headers", None):
                        page.calc_headers[0] = SELECTION_HEADER
                    page.load_table()

                    # Always show the source column after re-import.  QTableWidget
                    # keeps its old horizontal scroll position, which previously
                    # made a prepared file look like an initial import because
                    # the first calculated column was simply off-screen.
                    def show_source_column():
                        try:
                            page.table.horizontalScrollBar().setValue(0)
                            if page.table.columnCount() > 0:
                                page.table.setColumnHidden(0, False)
                        except RuntimeError:
                            pass

                    QTimer.singleShot(0, show_source_column)
                    unavailable = int(payload.get("unavailable") or 0)
                    message = (
                        "Подготовленный файл Target Price импортирован. "
                        f"Источники выбраны: {payload['selected']}; "
                        f"без выбора: {payload['blank']}."
                    )
                    if unavailable:
                        message += (
                            f" Без рассчитанных вариантов себестоимости: "
                            f"{unavailable}."
                        )
                    page.show_message(message)

                pbi._start_page_task(
                    page,
                    title="Target Price: импорт источников",
                    work=work_selection,
                    on_finished=done_selection,
                    read_resources={"products", "suppliers", "fixed_costs"},
                    write_resources={f"temp_target_price:{imported_by}"},
                    intro="Target Price: импорт источников выполняется в фоне.",
                    use_progress=True,
                )
                return

            # Initial product import does not need to create/update the target
            # supplier.  Cost-source calculation is independent of the target
            # supplier, which is selected only for the final 'Расчет Target'.
            supplier_id = None
            if not page.ui.cbx_NewSupplier.isChecked():
                current_supplier_id = page.ui.cbo_SupplName.currentData()
                if current_supplier_id is not None:
                    supplier_id = int(current_supplier_id)

            def work_initial(progress):
                from app.db.db import SessionLocal
                from app.imports.target_price_importer import TargetPriceImporter
                from app.services.target_price_service import TargetPriceService

                try:
                    progress("Target Price: импортирую файл...")
                    rows = TargetPriceImporter().read_excel(file_path)
                    with SessionLocal() as session:
                        service = TargetPriceService(session)
                        service.import_rows(
                            rows=rows,
                            batch_id=batch_id,
                            imported_by=imported_by,
                            supplier_id=supplier_id,
                        )
                        matched = service.automatch_temp_rows(batch_id, imported_by)
                        session.commit()
                    return {"count": len(rows), "matched": matched}
                finally:
                    SessionLocal.remove()

            def done_initial(payload):
                page._showing_options = False
                page.load_table()
                page.show_message(
                    f"Данные импортированы: {payload['count']}; "
                    f"автоматически сопоставлено: {payload['matched']}"
                )

            pbi._start_page_task(
                page,
                title="Target Price: импорт и сопоставление",
                work=work_initial,
                on_finished=done_initial,
                read_resources={"products", "product_articles"},
                write_resources={f"temp_target_price:{imported_by}"},
                intro="Target Price: импорт выполняется в фоне. Можно работать в других вкладках.",
                use_progress=True,
            )

        def calculate_background_fast() -> None:
            # Keep the GUI part intentionally light.  Validation/product creation
            # and the whole cost calculation happen in the worker thread.
            page._commit_open_editors()
            try:
                supplier_name = (
                    clean_multi_spaces(page.ui.cbo_SupplName.currentText())
                    or clean_multi_spaces(page.ui.line_NewSupplier.text())
                    or "Selection"
                )
                safe_supplier_name = "".join(
                    ch if ch not in r'<>:\"/\\|?*' else "_"
                    for ch in supplier_name
                )
                default = (
                    f"TargetPriceCalc_{safe_supplier_name}_"
                    f"{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
                )
                manual_full_costs = copy.deepcopy(
                    page.collect_manual_full_costs_from_table()
                )
                supplier_price_age_months = page.get_supplier_price_age_months()
                save_path, _ = QFileDialog.getSaveFileName(
                    page,
                    "Сохранить расчет",
                    str(Path(__file__).resolve().parents[1] / default),
                    "Excel files (*.xlsx)",
                )
                if not save_path:
                    return
                batch_id = page.batch_id
                imported_by = page.imported_by
                pending_deletes = set(getattr(page, "_pending_deletes", set()) or set())
            except Exception as exc:
                page.show_error_message(str(exc))
                return

            def launch_calculation() -> None:
                def work(progress):
                    from app.db.db import SessionLocal
                    from app.exports.target_price_exporter import TargetPriceExporter
                    from app.services.product_matching_service import MissingPackTypeError
                    from app.services.target_price_service import TargetPriceService

                    try:
                        with SessionLocal() as session:
                            try:
                                progress("Target Price: выполняю расчет...")
                                _delete_pending(
                                    session,
                                    pending_deletes,
                                    batch_id,
                                    imported_by,
                                )
                                service = TargetPriceService(session)
                                service.run_calculation(
                                    batch_id,
                                    imported_by,
                                    manual_full_costs=manual_full_costs,
                                    supplier_price_age_months=supplier_price_age_months,
                                )
                                output = TargetPriceExporter(session).export_calculated(
                                    batch_id,
                                    imported_by,
                                    save_path,
                                )
                                session.commit()
                                return {"status": "success", "path": str(output)}
                            except MissingPackTypeError as exc:
                                session.rollback()
                                return {"status": "missing_pack", "error": exc}
                    finally:
                        SessionLocal.remove()

                def done(payload):
                    if payload.get("status") == "missing_pack":
                        from app.db.models import TempTargetPriceImport
                        from app.utils.pack_type_prompt import resolve_missing_pack_for_temp_rows

                        if resolve_missing_pack_for_temp_rows(
                            page,
                            payload.get("error"),
                            model=TempTargetPriceImport,
                            batch_id=batch_id,
                            imported_by=imported_by,
                        ):
                            page.load_table()
                            launch_calculation()
                        return

                    pending = getattr(page, "_pending_deletes", None)
                    if isinstance(pending, set):
                        pending.clear()
                    snapshots = getattr(page, "_deleted_row_snapshots", None)
                    if isinstance(snapshots, list):
                        snapshots.clear()
                    page._showing_options = True
                    page.load_table()
                    output_path = payload.get("path")
                    if output_path:
                        QDesktopServices.openUrl(QUrl.fromLocalFile(str(output_path)))
                    page.show_message("Расчет выполнен")

                pbi._start_page_task(
                    page,
                    title="Target Price: расчет себестоимости",
                    work=work,
                    on_finished=done,
                    read_resources={
                        "products",
                        "supplier_prices",
                        "product_stock",
                        "product_uc3_history",
                    },
                    write_resources={f"temp_target_price:{imported_by}"},
                    intro=(
                        "Target Price: расчет выполняется в фоне. "
                        "Можно работать в других вкладках."
                    ),
                    use_progress=True,
                )

            launch_calculation()

        def calculate_target_background() -> None:
            # Do not use page._showing_options here.  The selected-source
            # workbook is self-contained and may have been imported in a new
            # application session.  Validate the actual restored DB batch in
            # the worker instead.
            page._commit_open_editors()
            try:
                supplier_id = page.ensure_supplier()
                supplier_name = (
                    clean_multi_spaces(page.ui.cbo_SupplName.currentText())
                    or clean_multi_spaces(page.ui.line_NewSupplier.text())
                    or "NoName"
                )
                safe_supplier_name = "".join(
                    ch if ch not in r'<>:\"/\\|?*' else "_"
                    for ch in supplier_name
                )
                default = (
                    f"TargetPrice_{safe_supplier_name}_"
                    f"{datetime.now().strftime('%Y-%m-%d_%H-%M-%S')}.xlsx"
                )
                save_path, _ = QFileDialog.getSaveFileName(
                    page,
                    "Сохранить Target Price",
                    str(Path(__file__).resolve().parents[1] / default),
                    "Excel files (*.xlsx)",
                )
                if not save_path:
                    return

                currency = clean_multi_spaces(page.ui.cbo_Currency.currentText()).upper()
                if currency == "-":
                    currency = ""
                batch_id = page.batch_id
                imported_by = page.imported_by
                fx_rate = page.parse_decimal_field(page.ui.line_ExchangeRate, "Курс")
                transport = page.parse_decimal_field(page.ui.line_Transport, "Транспорт")
                reexport = page.parse_percent_field(page.ui.line_Reexport, "Реэкспорт")
                insurance = page.parse_percent_field(page.ui.line_Insurance, "Insurance %")
                fx_markup = page.parse_percent_field(page.ui.line_FXMarkup, "FX markup %")
                fx_markup_abs = page.parse_decimal_field(page.ui.line_FXMarkupAbs, "FX markup abs")
                has_customs = page.ui.cbo_Customs.currentText() == "да"
                via_novo = page.ui.cbo_viaNovo.currentText() == "через Ново"
                manual_full_costs = copy.deepcopy(
                    page.collect_manual_full_costs_from_table()
                )
                pending_deletes = set(
                    getattr(page, "_pending_deletes", set()) or set()
                )
            except Exception as exc:
                page.show_error_message(str(exc))
                return

            def work(progress):
                from app.db.db import SessionLocal
                from app.db.models import (
                    Product,
                    TempTargetPriceImport,
                    TempTargetPriceOption,
                )
                from app.exports.target_price_exporter import TargetPriceExporter
                from app.services.target_price_service import TargetPriceService

                try:
                    with SessionLocal() as session:
                        _delete_pending(
                            session, pending_deletes, batch_id, imported_by
                        )

                        rows = (
                            session.query(TempTargetPriceImport)
                            .filter(
                                TempTargetPriceImport.batch_id == batch_id,
                                TempTargetPriceImport.imported_by == imported_by,
                            )
                            .order_by(
                                TempTargetPriceImport.import_row_no.asc(),
                                TempTargetPriceImport.id.asc(),
                            )
                            .all()
                        )
                        if not rows:
                            raise ValueError(
                                "Нет импортированных строк Target Price. "
                                "Импортируйте подготовленный файл с выбранными источниками."
                            )

                        # One-option rows are allowed to auto-select exactly as
                        # in the normal calculation path.  For every other row
                        # a real user selection must already be present.
                        service = TargetPriceService(session)
                        service.ensure_single_options_selected(
                            batch_id, imported_by
                        )
                        session.flush()

                        row_ids = [int(row.id) for row in rows]
                        option_rows = (
                            session.query(
                                TempTargetPriceOption.temp_import_id,
                                TempTargetPriceOption.id,
                            )
                            .filter(
                                TempTargetPriceOption.batch_id == batch_id,
                                TempTargetPriceOption.imported_by == imported_by,
                                TempTargetPriceOption.temp_import_id.in_(row_ids),
                            )
                            .all()
                            if row_ids else []
                        )
                        option_ids_by_row: dict[int, set[int]] = {}
                        for temp_import_id, option_id in option_rows:
                            option_ids_by_row.setdefault(
                                int(temp_import_id), set()
                            ).add(int(option_id))

                        # Rows for which no cost option exists are legitimate:
                        # there is simply nothing from which Target can be
                        # calculated. They must not block all other products.
                        unavailable_rows = [
                            row for row in rows
                            if row.selected_product_id is not None
                            and not option_ids_by_row.get(int(row.id))
                        ]
                        unavailable_ids = {
                            int(row.id) for row in unavailable_rows
                        }

                        # A missing product is still a real data problem.
                        missing_product_rows = [
                            row for row in rows
                            if row.selected_product_id is None
                        ]

                        # If options exist, the user must choose one. Do not
                        # silently skip a row where a choice was possible.
                        missing_selection_rows = [
                            row for row in rows
                            if row.selected_product_id is not None
                            and option_ids_by_row.get(int(row.id))
                            and row.selected_option_id is None
                        ]
                        blocking_rows = (
                            missing_product_rows + missing_selection_rows
                        )
                        if blocking_rows:
                            product_ids = {
                                int(row.selected_product_id)
                                for row in blocking_rows
                                if row.selected_product_id is not None
                            }
                            products = {
                                int(product.id): product.name or ""
                                for product in (
                                    session.query(Product)
                                    .filter(Product.id.in_(product_ids))
                                    .all()
                                    if product_ids else []
                                )
                            }
                            labels = []
                            for row in blocking_rows[:12]:
                                name = (
                                    products.get(
                                        int(row.selected_product_id), ""
                                    )
                                    if row.selected_product_id is not None
                                    else ""
                                )
                                labels.append(
                                    f"{row.import_row_no or row.id}: "
                                    f"{name or row.product_name or row.supplier_article or '-'}"
                                )
                            suffix = (
                                f"; ещё {len(blocking_rows) - len(labels)}"
                                if len(blocking_rows) > len(labels) else ""
                            )
                            raise ValueError(
                                "Не выбран источник для расчета Target: "
                                + "; ".join(labels)
                                + suffix
                            )

                        if unavailable_ids:
                            # These temp rows are needed only until the final
                            # Target calculation. Removing them inside this
                            # transaction lets the existing calculation service
                            # process all valid rows unchanged. On any later
                            # failure the transaction is rolled back.
                            session.query(TempTargetPriceImport).filter(
                                TempTargetPriceImport.batch_id == batch_id,
                                TempTargetPriceImport.imported_by == imported_by,
                                TempTargetPriceImport.id.in_(unavailable_ids),
                            ).delete(synchronize_session=False)
                            session.flush()

                        option_ids = {
                            int(row.selected_option_id) for row in rows
                            if int(row.id) not in unavailable_ids
                            and row.selected_option_id is not None
                        }
                        found_option_ids = {
                            int(option_id)
                            for (option_id,) in (
                                session.query(TempTargetPriceOption.id)
                                .filter(
                                    TempTargetPriceOption.batch_id == batch_id,
                                    TempTargetPriceOption.imported_by == imported_by,
                                    TempTargetPriceOption.id.in_(option_ids),
                                )
                                .all()
                                if option_ids else []
                            )
                        }
                        missing_option_ids = option_ids - found_option_ids
                        if missing_option_ids:
                            raise ValueError(
                                "Для части строк не восстановлены варианты себестоимости. "
                                "Повторно импортируйте подготовленный файл выбора."
                            )

                        progress("Target Price: выполняю расчет...")
                        saved_count = service.save_target_calculations(
                            batch_id=batch_id,
                            imported_by=imported_by,
                            target_supplier_id=supplier_id,
                            currency_code=currency,
                            fx_rate=fx_rate,
                            transport=transport,
                            reexport=reexport,
                            insurance=insurance,
                            fx_markup=fx_markup,
                            fx_markup_abs=fx_markup_abs,
                            has_customs=has_customs,
                            via_novo=via_novo,
                            manual_full_costs=manual_full_costs,
                        )
                        output = TargetPriceExporter(session).export_final(
                            batch_id, imported_by, save_path
                        )
                        service.delete_temp_rows_for_user(imported_by)
                        session.commit()
                        return {
                            "path": str(output),
                            "saved": int(saved_count or 0),
                            "skipped": len(unavailable_ids),
                        }
                finally:
                    SessionLocal.remove()

            def done(payload):
                pending = getattr(page, "_pending_deletes", None)
                if isinstance(pending, set):
                    pending.clear()
                snapshots = getattr(page, "_deleted_row_snapshots", None)
                if isinstance(snapshots, list):
                    snapshots.clear()

                output_path = (
                    payload.get("path")
                    if isinstance(payload, dict)
                    else payload
                )
                saved_count = (
                    int(payload.get("saved") or 0)
                    if isinstance(payload, dict)
                    else 0
                )
                skipped_count = (
                    int(payload.get("skipped") or 0)
                    if isinstance(payload, dict)
                    else 0
                )

                if output_path:
                    QDesktopServices.openUrl(
                        QUrl.fromLocalFile(str(output_path))
                    )
                page.start_new_batch()
                page._showing_options = False

                message = f"Расчет Target выполнен. Рассчитано: {saved_count}."
                if skipped_count:
                    message += (
                        " Без вариантов себестоимости пропущено: "
                        f"{skipped_count}."
                    )
                page.show_message(message)

            pbi._start_page_task(
                page,
                title="Target Price: расчет Target",
                work=work,
                on_finished=done,
                read_resources={
                    "products",
                    "supplier_prices",
                    "product_stock",
                    "product_uc3_history",
                },
                write_resources={
                    "target_price_calculations",
                    f"temp_target_price:{imported_by}",
                },
                intro=(
                    "Target Price: расчет выполняется в фоне. "
                    "Можно работать в других вкладках."
                ),
                use_progress=True,
            )

        page._background_target_price_import_slot = import_selection_aware
        page._background_target_price_calc_slot = calculate_background_fast
        page._background_target_price_save_slot = calculate_target_background
        page.ui.btn_Import.clicked.connect(import_selection_aware)
        page.ui.btn_CalcCost.clicked.connect(calculate_background_fast)
        page.ui.btn_Save.clicked.connect(calculate_target_background)

    attach_target_prices_with_selection._target_source_import = True
    attach_target_prices_with_selection._target_source_original = current_attach
    pbi._attach_target_prices = attach_target_prices_with_selection

def _install_target_page_connection_hook() -> None:
    """Attach the selection-aware Target Price workflow on every page instance.

    TargetPricesPage connects its original Import/Calculate/Save handlers inside
    setup_connections().  Patching only page_background_integration._attach_target_prices
    is not sufficient when no external caller invokes that attach function.  Hook
    setup_connections itself so the background/selection-aware handlers always replace
    the original buttons immediately after the page wires them.
    """
    from app.page_functions.target_prices_page import TargetPricesPage
    from app.utils import page_background_integration as pbi

    current_setup = TargetPricesPage.setup_connections
    if getattr(current_setup, "_target_selection_connection_hook", False):
        return

    def setup_connections_with_target_selection(self) -> None:
        current_setup(self)
        pbi._attach_target_prices(self)

    setup_connections_with_target_selection._target_selection_connection_hook = True
    setup_connections_with_target_selection._target_selection_original = current_setup
    TargetPricesPage.setup_connections = setup_connections_with_target_selection


def install_target_price_selection_workflow() -> None:
    global _installed
    if _installed:
        return
    _installed = True

    _install_target_supplier_fields_behavior()
    _install_selection_aware_import()
    _install_target_page_connection_hook()
