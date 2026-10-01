from __future__ import annotations

"""Customer Costs background calculation stability fixes.

The calculated Customer Costs workbook is intentionally written with openpyxl
instead of native Excel COM. This prevents the background task from remaining
in ``running`` when an invisible Excel process gets stuck in Workbook.Close()
or Excel.Quit() after the workbook has already been saved.

The page also always owns the pending-delete containers expected by the common
background completion handler.
"""

from decimal import Decimal
from pathlib import Path

_installed = False


def _excel_value(value):
    if value is None:
        return None
    if isinstance(value, Decimal):
        if value == 0:
            return None
        return float(value)
    if isinstance(value, (int, float)) and value == 0:
        return None
    return value


def _install_safe_calculated_export() -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from app.exports.customer_cost_exporter import CustomerCostExporter

    current = CustomerCostExporter.export_calculated
    if getattr(current, "_customer_cost_no_com", False):
        return

    def export_calculated_no_com(
        self,
        batch_id: str,
        imported_by: str,
        file_path: str | Path,
    ) -> Path:
        target_path = Path(file_path)
        if target_path.suffix.lower() != ".xlsx":
            target_path = target_path.with_suffix(".xlsx")
        target_path.parent.mkdir(parents=True, exist_ok=True)

        if target_path.exists():
            try:
                target_path.unlink()
            except PermissionError as exc:
                raise PermissionError(
                    f"Не удается перезаписать файл:\n{target_path}\n\n"
                    "Скорее всего, он открыт в Excel. Закрой файл и попробуй снова."
                ) from exc

        rows, max_opt = self._collect_export_rows(
            batch_id=batch_id,
            imported_by=imported_by,
        )

        base_headers = [
            "Our Product Name",
            "Дата",
            "Менеджер",
            "Клиент",
            "Код продукта",
            "Название продукта",
            "Фасовка",
            "Категория ABC",
            "Количество",
            "Объем л",
            "Вид закупки",
            "Условия оплаты",
            "Комментарии",
        ]

        dynamic_headers: list[str] = []
        for i in range(1, max_opt + 1):
            dynamic_headers.extend([
                f"Cost Novo with VAT_{i}",
                f"Full Cost Msk_{i}",
                f"Supplier_{i}",
                f"last update_{i}",
                f"FX rate_{i}",
                f"Currency_{i}",
            ])
        headers = base_headers + dynamic_headers

        wb = Workbook()
        ws = wb.active
        ws.title = "Calculated"
        ws.sheet_view.zoomScale = 85
        ws.freeze_panes = f"{get_column_letter(len(base_headers) + 1)}2"

        header_font = Font(name="Aptos Narrow", size=11, bold=True)
        body_font = Font(name="Aptos Narrow", size=11)
        header_alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        body_alignment = Alignment(vertical="center")
        base_fill = PatternFill("solid", fgColor="CDCDCD")
        cost_fill = PatternFill("solid", fgColor="00B0F0")
        source_fill = PatternFill("solid", fgColor="92D050")

        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(1, col_idx, header)
            cell.font = header_font
            cell.alignment = header_alignment
            if col_idx <= len(base_headers):
                cell.fill = base_fill
            else:
                offset = (col_idx - len(base_headers) - 1) % 6
                cell.fill = cost_fill if offset in {0, 1} else source_fill

        for row_idx, row in enumerate(rows, start=2):
            for col_idx, header in enumerate(headers, start=1):
                cell = ws.cell(row_idx, col_idx, _excel_value(row.get(header)))
                cell.font = body_font
                cell.alignment = body_alignment

        text_headers = {
            "Менеджер",
            "Клиент",
            "Код продукта",
            "Категория ABC",
        }
        integer_headers = {"Количество", "Объем л"}

        for col_idx, header in enumerate(headers, start=1):
            if header == "Дата" or header.startswith("last update_"):
                number_format = "dd.mm.yy;@"
            elif header in integer_headers:
                number_format = '# ##0;[Red]-# ##0;"-"'
            elif (
                header in text_headers
                or header.startswith("Supplier_")
                or header.startswith("Currency_")
            ):
                number_format = "@"
            elif (
                header.startswith("Cost Novo with VAT_")
                or header.startswith("Full Cost Msk_")
            ):
                number_format = "# ##0 ₽"
            elif header.startswith("FX rate_"):
                number_format = "# ##0"
            else:
                number_format = None

            if number_format:
                for row_idx in range(2, len(rows) + 2):
                    ws.cell(row_idx, col_idx).number_format = number_format

        base_widths = {
            "Our Product Name": 35.00,
            "Дата": 11.00,
            "Менеджер": 16.14,
            "Клиент": 18.14,
            "Код продукта": 16.00,
            "Название продукта": 31.14,
            "Фасовка": 10.50,
            "Категория ABC": 12.00,
            "Количество": 10.50,
            "Объем л": 10.50,
            "Вид закупки": 18.00,
            "Условия оплаты": 18.00,
            "Комментарии": 24.00,
        }
        for col_idx, header in enumerate(headers, start=1):
            width = base_widths.get(header)
            if width is not None:
                ws.column_dimensions[get_column_letter(col_idx)].width = width

        start_col = len(base_headers) + 1
        for _ in range(1, max_opt + 1):
            for offset, width in enumerate(
                (9.0, 9.0, 16.14, 10.14, 7.29, 9.14)
            ):
                ws.column_dimensions[
                    get_column_letter(start_col + offset)
                ].width = width
            start_col += 6

        last_col = get_column_letter(max(len(headers), 1))
        last_row = max(len(rows) + 1, 1)
        ws.auto_filter.ref = f"A1:{last_col}{last_row}"
        ws.row_dimensions[1].height = 30

        wb.save(target_path)
        return target_path

    export_calculated_no_com._customer_cost_no_com = True
    export_calculated_no_com._customer_cost_original = current
    CustomerCostExporter.export_calculated = export_calculated_no_com


def _install_page_state_guard() -> None:
    from app.utils import page_background_integration as pbi

    current = pbi._attach_customer_costs
    if getattr(current, "_customer_cost_state_guard", False):
        return

    original = current

    def attach_customer_costs_safe(page) -> None:
        # Common completion code clears these containers even when the user
        # never deleted rows via the context menu.
        if not hasattr(page, "_pending_deletes"):
            page._pending_deletes = set()
        if not hasattr(page, "_deleted_row_snapshots"):
            page._deleted_row_snapshots = []

        original(page)

    attach_customer_costs_safe._customer_cost_state_guard = True
    attach_customer_costs_safe._customer_cost_original = original
    pbi._attach_customer_costs = attach_customer_costs_safe


def install_customer_cost_background_fixes() -> None:
    global _installed
    if _installed:
        return

    _install_safe_calculated_export()
    _install_page_state_guard()
    _installed = True
