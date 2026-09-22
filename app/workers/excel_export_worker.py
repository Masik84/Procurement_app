from __future__ import annotations

from types import MethodType
from typing import Any, Callable

from PySide6.QtCore import QObject, Signal, Slot
from PySide6.QtWidgets import QAbstractButton, QTableWidget, QWidget

from app.utils.background_tasks import TaskConflictError, get_background_task_manager


def _set_owner_export_busy(owner: QObject, busy: bool) -> None:
    """Protect only the page that owns an Excel/calculation task.

    Other tabs stay interactive. The current page cannot start another save,
    import, reset or edit its table while the worker uses the same batch.
    """
    if not isinstance(owner, QWidget):
        return
    if busy:
        if getattr(owner, "_excel_export_busy_widgets", None):
            return
        saved: list[tuple[object, bool]] = []
        seen: set[int] = set()
        ui = getattr(owner, "ui", None)
        roots = [ui] if isinstance(ui, QWidget) else [owner]
        for root in roots:
            for widget in root.findChildren(QAbstractButton):
                if id(widget) in seen:
                    continue
                seen.add(id(widget))
                try:
                    saved.append((widget, bool(widget.isEnabled())))
                    widget.setEnabled(False)
                except RuntimeError:
                    pass
        table = getattr(owner, "table", None)
        if isinstance(table, QTableWidget) and id(table) not in seen:
            try:
                saved.append((table, bool(table.isEnabled())))
                table.setEnabled(False)
            except RuntimeError:
                pass
        owner._excel_export_busy_widgets = saved
    else:
        saved = list(getattr(owner, "_excel_export_busy_widgets", []) or [])
        owner._excel_export_busy_widgets = []
        for widget, enabled in saved:
            try:
                widget.setEnabled(bool(enabled))
            except RuntimeError:
                pass


def _install_safe_product_articles_export(owner: QObject) -> None:
    """Use openpyxl for Product Articles export instead of Excel COM.

    The Product Articles file is a simple four-column workbook. Using COM here
    added no functional value and was the place where a hidden Excel process
    could remain inside Quit() after the workbook had already been saved.
    """
    if owner.__class__.__name__ != "ProductArticlesPage":
        return
    if getattr(owner, "_safe_product_articles_export_installed", False):
        return

    def export_without_com(self, file_path: str, rows: list[dict]) -> None:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.sheet_view.zoomScale = 80

        headers = ["ID", "Product name", "Article", "Product name (variant)"]
        ws.append(headers)

        for row in rows:
            article = row.get("article", "")
            ws.append([
                row.get("id"),
                row.get("product_name", "") or "",
                "" if article is None else str(article),
                row.get("variant_name", "") or "",
            ])

        body_font = Font(name="Aptos Narrow", size=11)
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                cell.font = body_font

        header_fill = PatternFill(fill_type="solid", fgColor="CDCDCD")
        header_alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        for cell in ws[1]:
            cell.font = Font(name="Aptos Narrow", size=11, bold=True)
            cell.fill = header_fill
            cell.alignment = header_alignment

        ws.column_dimensions["A"].width = 10
        ws.column_dimensions["B"].width = 34
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 34

        # Article must remain text so leading zeroes are not lost.
        for cell in ws["C"]:
            cell.number_format = "@"

        ws.auto_filter.ref = ws.dimensions
        wb.save(file_path)
        wb.close()

    owner._export_product_articles_to_excel = MethodType(export_without_com, owner)
    owner._safe_product_articles_export_installed = True


class ExcelExportWorker(QObject):
    """Compatibility worker for legacy pages that still create their own QThread."""

    finished = Signal(object)
    error = Signal(str)

    def __init__(self, export_func: Callable[..., Any], *args, **kwargs) -> None:
        super().__init__()
        self._export_func = export_func
        self._args = args
        self._kwargs = kwargs

    @Slot()
    def run(self) -> None:
        try:
            self.finished.emit(self._export_func(*self._args, **self._kwargs))
        except Exception as exc:
            self.error.emit(str(exc).strip() or exc.__class__.__name__)


class _ExcelExportCallbackProxy(QObject):
    def __init__(
        self,
        *,
        on_finished: Callable[[Any], None] | None,
        on_error: Callable[[str], None] | None,
        finish_ui: Callable[[], None],
        parent: QObject | None = None,
    ) -> None:
        super().__init__(parent)
        self._on_finished = on_finished
        self._on_error = on_error
        self._finish_ui = finish_ui

    @Slot(object)
    def handle_finished(self, result: Any) -> None:
        self._finish_ui()
        if self._on_finished is not None:
            self._on_finished(result)

    @Slot(str)
    def handle_error(self, error_text: str) -> None:
        self._finish_ui()
        if self._on_error is not None:
            self._on_error(error_text)


def start_excel_export(
    owner: QObject,
    export_func: Callable[..., Any],
    *,
    on_finished: Callable[[Any], None] | None = None,
    on_error: Callable[[str], None] | None = None,
    button: Any | None = None,
    busy_text: str = "Формируется...",
    restore_text: str | None = None,
    args: tuple[Any, ...] = (),
    kwargs: dict[str, Any] | None = None,
) -> bool:
    """Run an Excel export through the application-wide background manager."""

    if getattr(owner, "_excel_export_task_handle", None) is not None:
        return False

    # Product Articles no longer needs a native Excel process at all.
    _install_safe_product_articles_export(owner)

    kwargs = kwargs or {}
    _set_owner_export_busy(owner, True)
    if button is not None:
        if restore_text is None:
            restore_text = button.text()
        button.setEnabled(False)
        button.setText(busy_text)

    def finish_ui() -> None:
        if button is not None:
            button.setEnabled(True)
            button.setText(restore_text or "Export Excel")
        setattr(owner, "_excel_export_task_handle", None)
        setattr(owner, "_excel_export_callback_proxy", None)
        setattr(owner, "_excel_export_thread", None)
        setattr(owner, "_excel_export_worker", None)
        _set_owner_export_busy(owner, False)

    proxy = _ExcelExportCallbackProxy(
        on_finished=on_finished,
        on_error=on_error,
        finish_ui=finish_ui,
        parent=owner,
    )
    setattr(owner, "_excel_export_callback_proxy", proxy)

    def work():
        return export_func(*args, **kwargs)

    owner_class = owner.__class__.__name__
    read_resources: set[str] = set()
    write_resources: set[str] = set()
    if owner_class == "SupplierPricesPage":
        read_resources = {
            "products",
            "supplier_price_calculations",
            "product_stock",
            "product_uc3_history",
        }
    elif owner_class == "CustomerCostsPage":
        read_resources = {"products", "supplier_prices", "product_stock"}
        write_resources = {"customer_price_calculations", "temp_customer_cost"}
    elif owner_class == "TargetPricesPage":
        read_resources = {
            "products", "supplier_prices", "product_stock", "product_uc3_history"
        }
        write_resources = {"target_price_calculations", "temp_target_price"}
    elif owner_class in {"PriceReportsPage", "CustomerCostsReportsPage", "OrderPlanningPage"}:
        read_resources = {"products", "product_stock", "supplier_prices"}

    task_title = {
        "SupplierPricesPage": "Прайс поставщика: Excel",
        "CustomerCostsPage": "Стоимость клиенту: Excel",
        "TargetPricesPage": "Target Price: Excel",
        "ProductSearchPage": "Поиск продуктов: Excel",
        "OrderPlanningPage": "Планирование закупок: Excel",
        "PriceReportsPage": "Отчет по ценам: Excel",
        "CustomerCostsReportsPage": "Отчет стоимости клиенту: Excel",
        "SuppliersPage": "Поставщики: Excel",
    }.get(owner_class, "Формирование Excel")

    try:
        handle = get_background_task_manager().start_task(
            task_title,
            work,
            read_resources=read_resources,
            write_resources=write_resources,
            owner=owner,
            intro=f"{task_title}: задача запущена.",
        )
    except TaskConflictError as exc:
        finish_ui()
        if on_error is not None:
            on_error(str(exc))
        return False
    except Exception as exc:
        finish_ui()
        if on_error is not None:
            on_error(str(exc))
        return False

    setattr(owner, "_excel_export_task_handle", handle)
    handle.finished.connect(proxy.handle_finished)
    handle.failed.connect(proxy.handle_error)
    return True
